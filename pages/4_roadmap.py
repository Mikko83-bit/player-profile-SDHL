from io import BytesIO
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def create_player_ppt(player_name, season_goal):
    prs = Presentation()
    # 16:9 Kuvasuhde
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # --- 1. OTSIKKO JA OIKEAN YLÄKULMAN TEKSTIT ---
    # Nimi (Vasen yläkulma)
    txBox = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.4), Inches(5.0), Inches(1.2)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Name:"
    p1.font.size = Pt(22)
    p1.font.bold = False

    p2 = tf.add_paragraph()
    p2.text = player_name
    p2.font.size = Pt(26)
    p2.font.bold = True

    # Athletes reflections (Oikea yläkulma)
    txBox_refl = slide.shapes.add_textbox(
        Inches(8.5), Inches(0.8), Inches(4.5), Inches(1.5)
    )
    tf_refl = txBox_refl.text_frame
    tf_refl.word_wrap = True

    p_r1 = tf_refl.paragraphs[0]
    p_r1.text = "Athletes reflections"
    p_r1.font.size = Pt(14)
    p_r1.font.underline = True

    p_r2 = tf_refl.add_paragraph()
    p_r2.text = "Dream goals:"
    p_r2.font.size = Pt(14)

    p_r3 = tf_refl.add_paragraph()
    p_r3.text = "Season:"
    p_r3.font.size = Pt(14)

    # --- LAATIKON LUONTIFUNKTIO ---
    def add_red_box(left, top, width, height, title, content_list):
        shape = slide.shapes.add_shape(1, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(179, 0, 0)  # Punainen
        shape.line.color.rgb = RGBColor(179, 0, 0)

        tf = shape.text_frame
        tf.word_wrap = True

        # Otsikko (Keltainen)
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 204, 0)
        p_title.alignment = PP_ALIGN.CENTER

        # Sisältö (Valkoinen)
        for item in content_list:
            p_item = tf.add_paragraph()
            p_item.text = str(item)
            p_item.font.size = Pt(13)
            p_item.font.color.rgb = RGBColor(255, 255, 255)
            p_item.alignment = PP_ALIGN.CENTER

    # --- 2. DREAM GOALS (Keskellä ylhäällä) ---
    add_red_box(
        Inches(4.1),
        Inches(0.8),
        Inches(3.8),
        Inches(1.5),
        "Dream goals",
        ["Career:", f"Season: {season_goal}"],
    )

    # --- 3. GAME ZONES ---
    add_red_box(
        Inches(1.8),
        Inches(2.7),
        Inches(3.6),
        Inches(1.3),
        "Offensive zone",
        ["Efficiency"],
    )

    add_red_box(
        Inches(5.8),
        Inches(2.7),
        Inches(3.8),
        Inches(1.3),
        "Defensive zone",
        ["Defending in own zone", "(comprehensive)"],
    )

    # --- 4. WEEKLY PROCESS & MENTAL ---
    add_red_box(
        Inches(1.2),
        Inches(4.7),
        Inches(3.4),
        "Mental",
        ["Emotional control", "Handling failure"],
    )

    add_red_box(
        Inches(6.5),
        Inches(4.7),
        Inches(3.6),
        Inches(1.3),
        "6x ice+off-ice training",
        ["1-2 games per week"],
    )

    # --- 5. POHJALAATIKOT (BASE) ---
    add_red_box(
        Inches(0.8),
        Inches(6.3),
        Inches(2.6),
        Inches(1.0),
        "Technical",
        ["Skating versatility"],
    )

    add_red_box(
        Inches(3.7),
        Inches(6.3),
        Inches(2.6),
        Inches(1.0),
        "Tactical",
        ["Special teams play"],
    )

    add_red_box(
        Inches(6.6),
        Inches(6.3),
        Inches(2.6),
        Inches(1.0),
        "Off-ice",
        ["Mobility & Speed"],
    )

    add_red_box(
        Inches(9.5),
        Inches(6.3),
        Inches(2.8),
        Inches(1.0),
        "Training (daily)",
        ["Daily theme"],
    )

    binary_output = BytesIO()
    prs.save(binary_output)
    binary_output.seek(0)
    return binary_output
